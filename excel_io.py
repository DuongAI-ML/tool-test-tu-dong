"""Đọc file Excel test case / test scenario của tester và tạo file mẫu trống.

Quy ước đọc sheet Test Case: MỖI BƯỚC MỘT DÒNG. Dòng có TC ID = bắt đầu TC mới (thông tin chung của TC ghi ở dòng
này); các dòng tiếp theo để trống TC ID (hoặc ô TC ID gộp/merge) = các bước tiếp theo của TC đó.
"""
import datetime as dt
import os
import re

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import columns as C
import keywords as K

LEVEL_STEP_KEYS = {c["key"] for c in C.TESTCASE_COLUMNS if c["level"] == "step"}
MODE_RANK = {"Tự động": 0, "Bán tự động": 1, "Tay": 2}


class ExcelError(Exception):
    pass


# ----------------------------------------------------------------------------- tiện ích
def cell_text(v):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, dt.datetime):
        return v.strftime("%d/%m/%Y") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%d/%m/%Y %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%d/%m/%Y")
    return str(v).strip()


def sheet_kind_by_name(name):
    n = C.norm(name)
    if n.startswith("_"):
        return None
    if "scenario" in n or "kich ban" in n or n in ("ts",):
        return "scenario"
    if "case" in n or n in ("tc", "testcase") or "test case" in n:
        return "testcase"
    if n in ("env", "environment", "moi truong", "bien", "cau hinh", "config") or "moi truong" in n or "bien" == n:
        return "env"
    if "element" in n or "doi tuong" in n or "locator" in n or "selector" in n:
        return "elements"
    if "huong dan" in n or "guide" in n or "tu khoa" in n or "danh muc" in n:
        return "skip"
    return None


def match_header(text, kind, profile=None):
    """Tên cột trong file -> key cột chuẩn (hoặc None)."""
    n = C.norm(text)
    if not n:
        return None, 0
    cols = C.CATALOG[kind]
    custom = {}
    if profile:
        for x in profile.get(kind, []):
            custom[C.norm(x.get("label"))] = x["key"]
    if n in custom:
        return custom[n], 3
    for c in cols:
        if n == C.norm(c["label"]) or n in c["syn"]:
            return c["key"], 2
    best, blen = None, 0
    for c in cols:
        for s in [C.norm(c["label"])] + c["syn"]:
            if len(s) >= 6 and (n.startswith(s + " ") or n.endswith(" " + s) or (" " + s + " ") in (" " + n + " ")):
                if len(s) > blen:
                    best, blen = c["key"], len(s)
    return (best, 1) if best else (None, 0)


def find_header(ws, kind, profile=None, max_scan=15):
    """Tìm dòng tiêu đề (dòng có nhiều tên cột khớp nhất) trong 15 dòng đầu."""
    best_row, best_score, best_map = None, 0, None
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row or 1, max_scan), values_only=True))
    for ri, row in enumerate(rows, start=1):
        mapping, used, score = [], set(), 0
        for ci, v in enumerate(row):
            t = cell_text(v)
            key, q = match_header(t, kind, profile) if t else (None, 0)
            if key and key in used:
                key = None
            if key:
                used.add(key)
                score += q
            mapping.append({"col": ci, "header": t, "key": key})
        if score > best_score:
            best_row, best_score, best_map = ri, score, mapping
    if not best_row or best_score < 3:
        return None, None
    return best_row, [m for m in best_map if m["header"]]


def guess_kind(ws, profile=None):
    r_tc, m_tc = find_header(ws, "testcase", profile)
    keys = {m["key"] for m in (m_tc or []) if m["key"]}
    if keys & {"buoc", "expected", "hanh_dong", "step_no"} and ("tc_id" in keys or "tieu_de" in keys):
        return "testcase"
    r_ts, m_ts = find_header(ws, "scenario", profile)
    keys2 = {m["key"] for m in (m_ts or []) if m["key"]}
    if "scenario" in keys2 and "ts_id" in keys2:
        return "scenario"
    return None


def parse_test_data(text):
    """'email=a@b.vn; pass=123' (hoặc mỗi dòng 1 cặp) -> {'email': 'a@b.vn', 'pass': '123'}"""
    out = {}
    for part in re.split(r"[;\n]+", text or ""):
        if "=" in part:
            k, v = part.split("=", 1)
            k = k.strip().strip("$").strip("{}").strip()
            if k and re.match(r"^[\wÀ-ỹ]+$", k):
                out[k] = v.strip()
        elif ":" in part:
            k, v = part.split(":", 1)
            k = k.strip()
            if k and re.match(r"^[\wÀ-ỹ]+$", k) and " " not in k:
                out[k] = v.strip()
    return out


def parse_sizes(text):
    """'1920x1080, 1366x768' / 'mobile, tablet' -> [(1920,1080), (1366,768)]"""
    out = []
    for part in re.split(r"[,;\n]+", text or ""):
        s = size_of(part)
        if s and s not in out:
            out.append(s)
    return out


PRESET_SIZES = {"mobile": (375, 812), "dien thoai": (375, 812), "iphone": (390, 844), "tablet": (768, 1024),
                "ipad": (768, 1024), "may tinh bang": (768, 1024), "laptop": (1366, 768),
                "desktop": (1920, 1080), "may tinh": (1920, 1080), "pc": (1920, 1080)}


def size_of(text):
    n = C.norm(text)
    if not n:
        return None
    if n in PRESET_SIZES:
        return PRESET_SIZES[n]
    m = re.search(r"(\d{3,4})\s*[x\*×X ]\s*(\d{3,4})", str(text))
    if m:
        w, h = int(m.group(1)), int(m.group(2))
        if 200 <= w <= 4000 and 200 <= h <= 4000:
            return (w, h)
    return None


def parse_browsers(text):
    out = []
    for part in re.split(r"[,;\n/]+", text or ""):
        n = C.norm(part)
        b = None
        if "edge" in n:
            b = "edge"
        elif "chrome" in n or "coc coc" in n:
            b = "chrome"
        elif "firefox" in n:
            b = "firefox"
        if b and b not in out:
            out.append(b)
    return out


# ----------------------------------------------------------------------------- đọc file
def read_workbook(path, profile=None):
    if not os.path.isfile(path):
        raise ExcelError("Không thấy file: %s" % path)
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        raise ExcelError("File đang bị khóa (có thể đang mở bằng chương trình khác). Hãy đóng file rồi thử lại.")
    except Exception as e:
        raise ExcelError("Không đọc được file Excel (chỉ hỗ trợ .xlsx/.xlsm): %s" % e)

    result = {"file": os.path.abspath(path), "file_name": os.path.basename(path),
              "project": os.path.splitext(os.path.basename(path))[0], "mtime": os.path.getmtime(path),
              "scenarios": [], "testcases": [], "envs": {}, "env_names": [], "elements": {},
              "headers": {}, "sheets": {}, "warnings": []}
    kinds = {}
    for ws in wb.worksheets:
        if ws.sheet_state != "visible" and ws.title.startswith("_"):
            continue
        k = sheet_kind_by_name(ws.title)
        if k == "skip":
            continue
        if k is None:
            k = guess_kind(ws, profile)
        if k and k not in kinds:
            kinds[k] = ws
    if "testcase" not in kinds and "scenario" not in kinds:
        raise ExcelError("Không tìm thấy sheet Test Case hoặc Test Scenario. Tên sheet nên có chữ 'Test Case' / "
                         "'Scenario', và dòng tiêu đề có các cột như TC ID, Bước thực hiện, Expected result…")
    if "scenario" in kinds:
        _read_scenarios(kinds["scenario"], result, profile)
    if "testcase" in kinds:
        _read_testcases(kinds["testcase"], result, profile)
    if "env" in kinds:
        _read_env(kinds["env"], result)
    if "elements" in kinds:
        _read_elements(kinds["elements"], result)
    if not result["env_names"]:
        result["env_names"] = ["Mặc định"]
        result["envs"] = {"Mặc định": {}}
    _link(result)
    for tc in result["testcases"]:
        analyse_tc(tc, result)
    return result


def _row_values(ws, header_row):
    for r_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        yield r_idx, [cell_text(v) for v in row]


def _read_scenarios(ws, result, profile):
    hr, mapping = find_header(ws, "scenario", profile)
    if not hr:
        result["warnings"].append("Sheet '%s': không nhận ra dòng tiêu đề Scenario." % ws.title)
        return
    result["headers"]["scenario"] = mapping
    result["sheets"]["scenario"] = {"name": ws.title, "header_row": hr}
    seen = set()
    for r, vals in _row_values(ws, hr):
        fields, extra = {}, {}
        for m in mapping:
            v = vals[m["col"]] if m["col"] < len(vals) else ""
            if not v:
                continue
            if m["key"]:
                fields[m["key"]] = v
            else:
                extra[m["header"]] = v
        if not fields and not extra:
            continue
        sid = fields.get("ts_id", "")
        if sid and len(fields) == 1 and not extra and " " in sid:
            continue        # dòng nhóm gộp ô, VD "LUỒNG 1 – TRUY CẬP, GIAO DIỆN & PHÂN QUYỀN (MH_01 – LU_01)"
        if not sid:
            if fields.get("scenario"):
                result["warnings"].append("Sheet '%s' dòng %d: scenario không có ID → bỏ qua." % (ws.title, r))
            continue
        if sid in seen:
            result["warnings"].append("Scenario ID '%s' bị trùng (dòng %d)." % (sid, r))
        seen.add(sid)
        result["scenarios"].append({"id": sid, "fields": fields, "extra": extra, "row": r})


def _read_testcases(ws, result, profile):
    hr, mapping = find_header(ws, "testcase", profile)
    if not hr:
        result["warnings"].append("Sheet '%s': không nhận ra dòng tiêu đề Test Case." % ws.title)
        return
    keys = {m["key"] for m in mapping if m["key"]}
    if "tc_id" not in keys:
        raise ExcelError("Sheet '%s' không có cột TC ID (mã test case). Hãy thêm cột 'TC ID'." % ws.title)
    result["headers"]["testcase"] = mapping
    result["sheets"]["testcase"] = {"name": ws.title, "header_row": hr}
    cur, ids, section = None, {}, ""
    for r, vals in _row_values(ws, hr):
        tc_f, st_f, extra_tc, extra_st = {}, {}, {}, {}
        for m in mapping:
            v = vals[m["col"]] if m["col"] < len(vals) else ""
            if not v:
                continue
            if m["key"] in LEVEL_STEP_KEYS:
                st_f[m["key"]] = v
            elif m["key"]:
                tc_f[m["key"]] = v
            else:
                extra_st[m["header"]] = v
        if not tc_f and not st_f and not extra_st:
            continue
        if tc_f.get("tc_id") and len(tc_f) == 1 and not st_f and not extra_st and is_section(tc_f["tc_id"]):
            section = tc_f["tc_id"]          # dòng tiêu đề nhóm, VD "LUỒNG NGHIỆP VỤ"
            continue
        if tc_f.get("tc_id"):
            tid = tc_f["tc_id"]
            n = ids.get(tid, 0) + 1
            ids[tid] = n
            key = tid if n == 1 else "%s#%d" % (tid, n)
            if n > 1:
                result["warnings"].append("TC ID '%s' bị trùng (dòng %d) → đặt tên nội bộ '%s'." % (tid, r, key))
            cur = {"key": key, "id": tid, "fields": tc_f, "extra": extra_st, "row": r, "steps": [], "section": section}
            result["testcases"].append(cur)
            extra_st = {}
        elif cur is None:
            result["warnings"].append("Sheet '%s' dòng %d: có dữ liệu nhưng chưa có TC ID ở trên → bỏ qua."
                                      % (ws.title, r))
            continue
        else:
            for k, v in tc_f.items():           # thông tin chung ghi ở dòng sau (không gộp ô) vẫn nhận
                cur["fields"].setdefault(k, v)
        if st_f or extra_st:
            cur["steps"].append({"fields": st_f, "extra": extra_st, "row": r})
    for tc in result["testcases"]:
        split_steps(tc)
        for i, s in enumerate(tc["steps"], start=1):
            s["no"] = s["fields"].get("step_no") or str(i)


def is_section(text):
    """'LUỒNG NGHIỆP VỤ', 'Chức năng đăng ký' … (không có số, có khoảng trắng) = dòng nhóm, không phải TC ID."""
    t = (text or "").strip()
    return bool(t) and " " in t and not re.search(r"\d", t) and len(t) <= 80


NUM_RE = re.compile(r"(?m)^\s*(?:bước\s*)?(\d{1,2})\s*[.)/:-]\s*", re.I)


def split_numbered(text):
    """'1. Nhập A\n2. Nhập B' -> [('1','Nhập A'), ('2','Nhập B')]; không đánh số -> []"""
    parts = NUM_RE.split(text or "")
    if len(parts) < 4:
        return []
    out = []
    for i in range(1, len(parts) - 1, 2):
        out.append((parts[i], re.sub(r"\s+", " ", parts[i + 1]).strip()))
    return out


def split_steps(tc):
    """File viết TẤT CẢ các bước trong 1 ô (đánh số 1. 2. 3.) → tách thành từng bước để chạy / ghi kết quả."""
    if len(tc["steps"]) != 1:
        return
    st = tc["steps"][0]
    f = st["fields"]
    steps = split_numbered(f.get("buoc", ""))
    if len(steps) < 2:
        return
    exp = dict(split_numbered(f.get("expected", "")))
    acts = {k: dict(split_numbered(f.get(k, ""))) for k in ("hanh_dong", "doi_tuong", "gia_tri")}
    new = []
    for i, (no, text) in enumerate(steps):
        nf = {"buoc": text, "step_no": no}
        if exp:
            if exp.get(no):
                nf["expected"] = exp[no]
        elif i == len(steps) - 1 and f.get("expected"):
            nf["expected"] = f["expected"]          # 1 kết quả mong đợi chung → gắn vào bước cuối
        for k, d in acts.items():
            if d.get(no):
                nf[k] = d[no]
        for k, v in f.items():
            if k not in ("buoc", "expected", "step_no", "hanh_dong", "doi_tuong", "gia_tri") and i == 0:
                nf[k] = v
        new.append({"fields": nf, "extra": st.get("extra", {}) if i == 0 else {}, "row": st.get("row")})
    tc["steps"] = new
    tc["steps_merged"] = True


def _read_env(ws, result):
    rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    if not rows:
        return
    head = rows[0]
    names = [h for h in head[1:] if h]
    if not names or (len(names) == 1 and C.norm(names[0]) in ("gia tri", "value")):
        names = ["Mặc định"]
    envs = {n: {} for n in names}
    for r in rows[1:]:
        var = (r[0] if r else "").strip().strip("$").strip("{}").strip()
        if not var:
            continue
        for i, n in enumerate(names, start=1):
            envs[n][var] = r[i] if i < len(r) else ""
    result["envs"] = envs
    result["env_names"] = names


def _read_elements(ws, result):
    rows = [[cell_text(v) for v in row] for row in ws.iter_rows(values_only=True)]
    rows = [r for r in rows if any(r)]
    for r in rows[1:] if rows and C.norm(rows[0][0]) in ("ten", "ten doi tuong", "name", "element", "doi tuong",
                                                          "ten element") else rows:
        if len(r) >= 2 and r[0] and r[1]:
            result["elements"][r[0].strip()] = r[1].strip()


def _link(result):
    by_id = {s["id"]: s for s in result["scenarios"]}
    for s in result["scenarios"]:
        s["tcs"] = []
    for tc in result["testcases"]:
        sid = tc["fields"].get("scenario_id", "")
        if sid and sid in by_id:
            by_id[sid]["tcs"].append(tc["key"])
        elif sid:
            tc.setdefault("warnings", []).append("Scenario ID '%s' không có trong sheet Scenario." % sid)


def analyse_tc(tc, result=None):
    """Xác định CÁCH TEST thực tế của TC + lý do, để báo rõ cho người dùng (không im lặng)."""
    warns = tc.setdefault("warnings", [])
    reasons = []
    steps_auto, steps_manual, steps_bad = [], [], []
    for s in tc["steps"]:
        raw = s["fields"].get("hanh_dong", "")
        name = K.lookup(raw)
        s["action"] = name or ""
        if name is None:
            s["kind"] = "invalid"
            steps_bad.append(s["no"])
            warns.append("Bước %s: hành động '%s' không có trong danh sách → bước này sẽ báo LỖI khi chạy."
                         % (s["no"], raw))
        elif name == "":
            s["kind"] = "manual"
            steps_manual.append(s["no"])
        elif name in K.MANUAL:
            s["kind"] = "manual"
            steps_manual.append(s["no"])
        else:
            s["kind"] = "auto"
            steps_auto.append(s["no"])
            if name == "chạy_tc" and result is not None:
                target = s["fields"].get("gia_tri", "")
                if not any(t["id"] == target or t["key"] == target for t in result["testcases"]):
                    warns.append("Bước %s: gọi TC '%s' nhưng không có TC này trong file." % (s["no"], target))
    if not tc["steps"]:
        steps_mode = "Tay"
        reasons.append("TC chưa có bước nào")
    elif not steps_auto and not steps_bad:
        steps_mode = "Tay"
        if any(not s["fields"].get("hanh_dong") for s in tc["steps"]):
            reasons.append("chưa có cột Hành động cho các bước")
        else:
            reasons.append("tất cả các bước là kiểm tra bằng mắt")
    elif steps_manual:
        steps_mode = "Bán tự động"
        reasons.append("bước %s cần người kiểm tra" % ", ".join(steps_manual))
    else:
        steps_mode = "Tự động"

    declared = C.mode_of(tc["fields"].get("cach_test", ""))
    if tc["fields"].get("cach_test") and not declared:
        warns.append("Cột Cách test ghi '%s' — không hiểu, tool tự xác định." % tc["fields"]["cach_test"])
    type_mode = C.TYPE_DEFAULT_MODE.get(C.norm(tc["fields"].get("loai_test", "")), "")
    wanted = declared or type_mode or "Tự động"
    mode = steps_mode if MODE_RANK[steps_mode] >= MODE_RANK[wanted] else wanted
    if declared and MODE_RANK[steps_mode] > MODE_RANK[declared]:
        warns.append("Cột Cách test ghi '%s' nhưng %s → tool chạy dạng %s." % (declared, "; ".join(reasons),
                                                                                steps_mode.upper()))
    if declared and declared == mode and mode != "Tự động":
        reasons.insert(0, "cột Cách test ghi '%s'" % declared)
    elif mode != steps_mode:
        if declared:
            reasons.append("cột Cách test ghi '%s'" % declared)
        else:
            reasons.append("loại test '%s' mặc định là %s" % (tc["fields"].get("loai_test"), wanted))
    tc["mode"] = mode
    tc["mode_declared"] = declared
    tc["review_at_end"] = mode == "Bán tự động" and not steps_manual and not steps_bad
    if tc["review_at_end"]:
        reasons.append("tool chạy hết các bước rồi DỪNG để bạn duyệt kết quả")
    tc["mode_reason"] = "; ".join(reasons)
    tc["manual_steps"] = steps_manual
    return tc


# ----------------------------------------------------------------------------- tạo file mẫu
THIN = Side(style="thin", color="BFC5CE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
RESULT_FILL = PatternFill("solid", fgColor="7F6000")
AUTO_FILL = PatternFill("solid", fgColor="375623")
HEAD_FONT = Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")
WIDTHS = {"tc_id": 16, "ts_id": 14, "scenario_id": 14, "tieu_de": 34, "scenario": 50, "buoc": 36, "expected": 34,
          "precondition": 26, "test_data": 24, "hanh_dong": 20, "doi_tuong": 24, "gia_tri": 24, "actual": 30,
          "mo_ta": 30, "ghi_chu": 24, "kich_thuoc": 20, "step_no": 7, "ac": 12, "nhom": 14, "module": 16,
          "loai_test": 14, "priority": 10, "cach_test": 13, "status": 11, "status_buoc": 11, "anh": 14}


def enabled_columns(profile, kind):
    out = []
    for x in profile.get(kind, []):
        if x.get("on") and x["key"] in C.BY_KEY[kind]:
            c = dict(C.BY_KEY[kind][x["key"]])
            c["label"] = x.get("label") or c["label"]
            out.append(c)
    return out


def style_header(ws, cols, row=1):
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=c["label"])
        cell.font = HEAD_FONT
        cell.fill = RESULT_FILL if c.get("result") else (AUTO_FILL if c.get("group") == "Tự động hóa" else HEAD_FILL)
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = BORDER
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(c["key"], 15)
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def _lists_sheet(wb):
    ws = wb.create_sheet("_DanhMuc")
    lists = {
        "priority": C.PRIORITIES, "severity": C.SEVERITIES, "cach_test": C.CACH_TEST,
        "loai_test": [t[0] for t in C.TEST_TYPES], "nhom": [t[0] for t in C.TEST_TYPES],
        "hanh_dong": [k["name"] for k in K.KEYWORDS],
        "status": ["Pass", "Fail", "Blocked", "Chưa test", "Skip"],
    }
    refs = {}
    for ci, (k, vals) in enumerate(lists.items(), start=1):
        ws.cell(row=1, column=ci, value=k)
        for ri, v in enumerate(vals, start=2):
            ws.cell(row=ri, column=ci, value=v)
        col = get_column_letter(ci)
        refs[k] = "'_DanhMuc'!$%s$2:$%s$%d" % (col, col, len(vals) + 1)
    refs["status_buoc"] = refs["status"]
    ws.sheet_state = "hidden"
    return refs


def add_validations(ws, cols, refs, first_row=2, last_row=2000):
    for i, c in enumerate(cols, start=1):
        if c["key"] in refs:
            # cột Nhóm/Loại test, Hành động: cho phép gõ giá trị khác danh sách (showErrorMessage=False)
            strict = c["key"] in ("priority", "cach_test", "severity")
            dv = DataValidation(type="list", formula1="=" + refs[c["key"]], allow_blank=True,
                                showErrorMessage=strict)
            if strict:
                dv.error = "Chọn giá trị trong danh sách"
            col = get_column_letter(i)
            dv.add("%s%d:%s%d" % (col, first_row, col, last_row))
            ws.add_data_validation(dv)


def write_guide_sheet(wb):
    ws = wb.create_sheet("HuongDan")
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 26
    ws.column_dimensions["F"].width = 26
    ws["A1"] = "HƯỚNG DẪN NHANH"
    ws["A1"].font = Font(bold=True, size=14)
    tips = [
        "• Sheet TestCase: MỖI BƯỚC MỘT DÒNG. Dòng có TC ID là dòng đầu của TC; các dòng bước sau để trống TC ID.",
        "• Cột Hành động / Đối tượng / Giá trị là phần tool chạy tự động. Bỏ trống Hành động = bước đó TEST TAY.",
        "• Đối tượng: ghi như người đọc (Ô Email, Nút \"Đăng nhập\", Link \"Quên mật khẩu\") hoặc tên trong sheet "
        "Elements, hoặc selector CSS/XPath. Thêm ' #2' để chọn phần tử thứ 2.",
        "• Biến: ${base_url}, ${user}… lấy từ sheet Env; ${ngau_nhien} = số ngẫu nhiên; ${hom_nay} = ngày hôm nay; "
        "cột Test data ghi 'ten=giá trị; ten2=giá trị' để tạo biến riêng cho TC.",
        "• Cột Kích thước màn hình: '1920x1080, 768x1024, 375x812' hoặc 'desktop, tablet, mobile' → chạy TC ở từng "
        "kích thước.",
        "• Cách test: Tự động / Bán tự động / Tay. Bán tự động & Tay: tool sẽ DỪNG và hỏi bạn Pass/Fail.",
    ]
    for i, t in enumerate(tips, start=2):
        ws.cell(row=i, column=1, value=t)
    r = len(tips) + 3
    for i, h in enumerate(["Hành động", "Đối tượng", "Giá trị", "Ý nghĩa", "Ví dụ đối tượng", "Ví dụ giá trị"],
                          start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = HEAD_FONT
        c.fill = HEAD_FILL
    for k in K.KEYWORDS:
        r += 1
        ws.cell(row=r, column=1, value=k["name"]).font = Font(bold=True, color="C00000" if k["manual"] else "000000")
        ws.cell(row=r, column=2, value=k["target"] or "-")
        ws.cell(row=r, column=3, value=k["value"] or "-")
        ws.cell(row=r, column=4, value=k["desc"]).alignment = WRAP
        ws.cell(row=r, column=5, value=k["ex_target"])
        ws.cell(row=r, column=6, value=k["ex_value"])
    r += 2
    ws.cell(row=r, column=1, value="LOẠI TEST").font = Font(bold=True, size=12)
    for name, mode, note in C.TEST_TYPES:
        r += 1
        ws.cell(row=r, column=1, value=name).font = Font(bold=True)
        ws.cell(row=r, column=2, value=mode)
        ws.cell(row=r, column=4, value=note)


def create_template(path, profile, sample=None):
    """Tạo file Excel trống theo các cột đang bật trong hồ sơ cột. sample = dữ liệu ví dụ (dict) nếu có."""
    wb = openpyxl.Workbook()
    refs = _lists_sheet(wb)
    ws_ts = wb.active
    ws_ts.title = "Scenario"
    ts_cols = [c for c in enabled_columns(profile, "scenario") if not c["result"]]
    style_header(ws_ts, ts_cols)
    add_validations(ws_ts, ts_cols, refs)
    ws_tc = wb.create_sheet("TestCase")
    tc_cols = [c for c in enabled_columns(profile, "testcase") if not c["result"]]
    for need in ("tc_id", "buoc", "hanh_dong", "doi_tuong", "gia_tri"):
        if need not in [c["key"] for c in tc_cols]:
            c = dict(C.BY_KEY["testcase"][need])
            tc_cols.append(c)
    style_header(ws_tc, tc_cols)
    add_validations(ws_tc, tc_cols, refs)
    ws_env = wb.create_sheet("Env")
    env_head = ["Biến"] + ((sample or {}).get("env_names") or ["Dev", "Staging"])
    for i, h in enumerate(env_head, start=1):
        c = ws_env.cell(row=1, column=i, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        ws_env.column_dimensions[get_column_letter(i)].width = 22 if i == 1 else 40
    env_rows = (sample or {}).get("env_rows") or [["base_url", "http://dev.congty.local", "http://staging.congty.local"],
                                                  ["user", "", ""], ["pass", "", ""]]
    for r, row in enumerate(env_rows, start=2):
        for i, v in enumerate(row, start=1):
            ws_env.cell(row=r, column=i, value=v)
    ws_el = wb.create_sheet("Elements")
    for i, h in enumerate(["Tên đối tượng", "Selector (CSS / XPath / text=…)", "Ghi chú"], start=1):
        c = ws_el.cell(row=1, column=i, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    ws_el.column_dimensions["A"].width = 28
    ws_el.column_dimensions["B"].width = 50
    ws_el.column_dimensions["C"].width = 40
    for r, row in enumerate((sample or {}).get("elements") or [], start=2):
        for i, v in enumerate(row, start=1):
            ws_el.cell(row=r, column=i, value=v)
    if sample:
        _fill_rows(ws_ts, ts_cols, sample.get("scenarios", []))
        _fill_rows(ws_tc, tc_cols, sample.get("testcases", []), group_key="tc_id")
    write_guide_sheet(wb)
    wb.move_sheet("_DanhMuc", offset=len(wb.sheetnames))
    wb.active = 1 if sample else 0
    try:
        wb.save(path)
    except PermissionError:
        raise ExcelError("Không ghi được file (đang mở trong Excel?): %s" % path)
    return path


def _fill_rows(ws, cols, rows, group_key=None):
    idx = {c["key"]: i for i, c in enumerate(cols, start=1)}
    band = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FA")]
    b = -1
    for r, row in enumerate(rows, start=2):
        if group_key is None or row.get(group_key):
            b += 1
        for k, i in idx.items():
            cell = ws.cell(row=r, column=i, value=row.get(k))
            cell.alignment = WRAP
            cell.border = BORDER
            if group_key:
                cell.fill = band[b % 2]
