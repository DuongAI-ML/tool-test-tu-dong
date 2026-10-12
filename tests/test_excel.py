"""Test đọc Excel theo nhiều kiểu file tester hay dùng. Chạy: python -m pytest tests -q  (hoặc python tests/test_excel.py)"""
import os
import sys
import tempfile

import openpyxl

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import columns as C  # noqa: E402
import excel_io as X  # noqa: E402
import report as R  # noqa: E402
import sample  # noqa: E402


def _tmp(name):
    return os.path.join(tempfile.mkdtemp(prefix="tt_"), name)


def test_scenario_cua_tester():
    """Đúng định dạng scenario tester gửi: ID | Nhóm | Test Scenario | AC liên quan | Priority, sheet tên lạ."""
    p = _tmp("ds.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["ID", "Nhóm", "Test Scenario", "AC liên quan", "Priority"])
    ws.append(["TS_DHTD_001", "usability", "Kiểm tra giao diện và bố cục mặc định của trang Đơn hàng tuyển dụng",
               "AC_01", "Medium"])
    ws.append(["TS_DHTD_002", "usability", "Kiểm tra khả năng hiển thị khi thay đổi kích thước màn hình", "AC_02",
               "High"])
    wb.save(p)
    d = X.read_workbook(p)
    assert [s["id"] for s in d["scenarios"]] == ["TS_DHTD_001", "TS_DHTD_002"]
    f = d["scenarios"][0]["fields"]
    assert f["nhom"] == "usability" and f["ac"] == "AC_01" and f["priority"] == "Medium"
    assert "bố cục" in f["scenario"]


def test_testcase_tieng_viet_o_gop_va_tieu_de_dong_3():
    """Tên cột tiếng Việt, dòng tiêu đề ở dòng 3, ô TC ID gộp (merge) qua các dòng bước."""
    p = _tmp("tc.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Case"
    ws.append(["DANH SÁCH TEST CASE — MODULE ĐƠN HÀNG"])
    ws.append([])
    ws.append(["Mã TC", "Tên test case", "Điều kiện tiên quyết", "STT", "Các bước thực hiện", "Kết quả mong đợi",
               "Kết quả thực tế", "Trạng thái", "Người test", "Cột riêng của tôi"])
    ws.append(["TC01", "Đăng nhập đúng", "Có tài khoản", 1, "Mở trang login", "Hiện form", None, None, None, "x1"])
    ws.append([None, None, None, 2, "Nhập email, mật khẩu", "", None, None, None, None])
    ws.append([None, None, None, 3, "Bấm Đăng nhập", "Vào trang chủ", None, None, None, None])
    ws.append(["TC02", "Đăng nhập sai", None, 1, "Nhập sai mật khẩu", "Báo lỗi", None, None, None, None])
    ws.merge_cells("A4:A6")
    ws.merge_cells("B4:B6")
    wb.save(p)
    d = X.read_workbook(p)
    assert [t["id"] for t in d["testcases"]] == ["TC01", "TC02"]
    t1 = d["testcases"][0]
    assert len(t1["steps"]) == 3
    assert t1["fields"]["precondition"] == "Có tài khoản"
    assert t1["steps"][2]["fields"]["expected"] == "Vào trang chủ"
    # không có cột Hành động → TEST TAY, phải có lý do rõ ràng
    assert t1["mode"] == "Tay" and "Hành động" in t1["mode_reason"]
    m = {x["header"]: x["key"] for x in d["headers"]["testcase"]}
    assert m["Kết quả thực tế"] == "actual" and m["Trạng thái"] == "status" and m["Cột riêng của tôi"] is None
    assert t1["extra"] == {"Cột riêng của tôi": "x1"}


def test_cach_test_va_canh_bao():
    p = _tmp("mode.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TestCase"
    ws.append(["TC ID", "Loại test", "Cách test", "Bước thực hiện", "Hành động", "Đối tượng", "Giá trị"])
    ws.append(["A", "Functional", "Tự động", "Mở", "mở_trang", "", "http://x"])
    ws.append([None, None, None, "Nhìn", "", "", ""])                  # bước không có hành động
    ws.append(["B", "Usability", "", "Mở", "open", "", "http://x"])      # usability -> Tay theo loại test
    ws.append(["C", "Functional", "", "Mở", "mo trang", "", "http://x"])  # không dấu vẫn nhận
    ws.append([None, None, None, "Sai", "bay_len_troi", "", ""])         # hành động không hợp lệ
    ws.append(["D", "UI", "", "Mở", "mở_trang", "", "http://x"])         # UI -> bán tự động, duyệt cuối
    wb.save(p)
    d = {t["id"]: t for t in X.read_workbook(p)["testcases"]}
    assert d["A"]["mode"] == "Bán tự động" and any("Tự động" in w for w in d["A"]["warnings"])
    assert d["B"]["mode"] == "Tay" and "Usability" in d["B"]["mode_reason"]
    assert d["C"]["steps"][0]["action"] == "mở_trang"
    assert d["C"]["steps"][1]["kind"] == "invalid" and d["C"]["warnings"]
    assert d["D"]["mode"] == "Bán tự động" and d["D"]["review_at_end"]


def test_file_mau_doc_lai_duoc():
    p = _tmp("mau.xlsx")
    prof = C.default_profile()
    X.create_template(p, prof, sample.sample_data())
    d = X.read_workbook(p)
    assert len(d["testcases"]) == len(set(r["tc_id"] for r in sample.TESTCASES if r.get("tc_id")))
    assert d["env_names"] == ["Demo", "Staging"] and d["envs"]["Demo"]["base_url"] == "${demo_url}"
    assert "Dòng trong bảng" in d["elements"]
    assert not d["warnings"]
    for t in d["testcases"]:
        assert not t["warnings"], (t["id"], t["warnings"])
    # file trống
    p2 = _tmp("trong.xlsx")
    X.create_template(p2, prof)
    wb = openpyxl.load_workbook(p2)
    assert {"Scenario", "TestCase", "Env", "Elements", "HuongDan", "_DanhMuc"} <= set(wb.sheetnames)


def test_ten_file_xuat():
    ctx = {"du_an": "Đơn hàng: TD/2026", "ngay": "2026-10-12", "gio": "1530", "moi_truong": "Staging",
           "lan_chay": "2", "nguoi_test": "Lan", "module": "X", "trinh_duyet": "edge"}
    assert R.build_name("KetQua_{du_an}_{ngay}_{gio}", ctx) == "KetQua_Đơn hàng_ TD_2026_2026-10-12_1530.xlsx"
    assert R.build_name("", ctx).startswith("KetQua_")
    assert R.build_name("BaoCao_{moi_truong}_lan{lan_chay}.xlsx", ctx) == "BaoCao_Staging_lan2.xlsx"
    assert R.build_name("CON", ctx) == "_CON.xlsx"
    assert R.build_name("a{khong_co}b", ctx) == "a{khong_co}b.xlsx"


def test_bien_va_kich_thuoc():
    assert X.parse_test_data("email=a@b.vn; sdt=0901\npass=1=2") == {"email": "a@b.vn", "sdt": "0901",
                                                                     "pass": "1=2"}
    assert X.parse_sizes("1920x1080, mobile; 768 x 1024, rác") == [(1920, 1080), (375, 812), (768, 1024)]
    assert X.parse_browsers("Edge, Chrome / Cốc Cốc") == ["edge", "chrome"]


def test_file_test_case_that_cua_tester():
    """Đúng mẫu ảnh tester gửi: các bước gộp trong 1 ô, có dòng nhóm 'LUỒNG NGHIỆP VỤ', cột Lần 1/2/3, Bug #."""
    p = _tmp("tc_that.xlsx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test case"
    ws.append(["Mã TC", "Chức năng", "Requirement / Scenario", "Tên TC", "Preconditions", "Environment", "Steps",
               "Kết quả mong đợi", "Kết quả thực tế", "Data ID", "Test data", "Priority", "Kết quả test (Lần 1)",
               "Kết quả test (Lần 2)", "Kết quả test (Lần 3)", "Trạng thái Test tổng", "Bug #"])
    ws.append(["LUỒNG NGHIỆP VỤ"])
    ws.append(["TC001", "Đăng ký tài khoản", "AC01 / Đăng ký hợp lệ", "Kiểm tra đăng ký thành công với MK 6 ký tự",
               "Trang Register hoạt động", "Chrome / Windows 11",
               "1. Nhập Họ, Họ tên, Địa chỉ email hợp lệ.\n\n2. Nhập MK đúng 6 ký tự.\n\n3. Nhập Mật khẩu (nhắc lại) "
               "trùng khớp..\n\n4. Tích vào ô \"Đồng ý\"\n\n5. Click vào ô \"Gửi đi\"",
               "Báo đăng ký thành công, tạo TK Pending, gửi mail.", "Lỗi: Hiện thị \"Lỗi\"", "TD-001",
               "Email: test01@gmail.com", "High", "FAIL", None, None, "FAIL", None])
    wb.save(p)
    d = X.read_workbook(p)
    assert [t["id"] for t in d["testcases"]] == ["TC001"], [t["id"] for t in d["testcases"]]
    t = d["testcases"][0]
    assert len(t["steps"]) == 5 and t["steps"][4]["fields"]["buoc"].startswith("Click vào ô")
    assert t["steps"][4]["fields"]["expected"].startswith("Báo đăng ký thành công")
    assert t["section"] == "LUỒNG NGHIỆP VỤ" and t["fields"]["data_id"] == "TD-001"
    m = {x["header"]: x["key"] for x in d["headers"]["testcase"]}
    assert m["Kết quả test (Lần 1)"] == "kq_lan1" and m["Trạng thái Test tổng"] == "status"
    assert m["Requirement / Scenario"] == "requirement" and m["Preconditions"] == "precondition"
    assert m["Environment"] == "moi_truong" and m["Steps"] == "buoc"
    d2 = X.read_workbook(p, C.company_profile())
    assert {x["header"]: x["key"] for x in d2["headers"]["testcase"]}["Bug #"] == "anh"


if __name__ == "__main__":
    ok = 0
    for name, fn in list(globals().items()):
        if name.startswith("test_") and callable(fn):
            fn()
            ok += 1
            print("PASS", name)
    print("%d test OK" % ok)
