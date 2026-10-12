"""Xuất kết quả chạy test: Excel (theo các cột đang tick), HTML, và bản sao file gốc có ghi kết quả."""
import datetime as dt
import html
import os
import re
import shutil

import copy

import openpyxl
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import columns as C
import excel_io as X

STATUS_FILL = {"Pass": "C6EFCE", "Fail": "FFC7CE", "Blocked": "FFE699", "Chưa test": "F4B183", "Skip": "D9D9D9"}
STATUS_FONT = {"Pass": "006100", "Fail": "9C0006", "Blocked": "7F6000", "Chưa test": "843C0C", "Skip": "595959"}
BAD_CHARS = r'[\\/:*?"<>|\r\n\t]'
VARS_HELP = {"du_an": "Tên dự án (tên file Excel)", "module": "Module (nếu chạy 1 module)", "ngay": "Ngày (2026-10-12)",
             "gio": "Giờ phút (1530)", "moi_truong": "Môi trường", "lan_chay": "Lần chạy thứ mấy trong ngày",
             "nguoi_test": "Tên người test", "trinh_duyet": "Trình duyệt"}


def name_context(data, settings, run_no=1):
    res = data.get("results", [])
    mods = sorted({r["fields"].get("module", "") for r in res if r["fields"].get("module")})
    started = dt.datetime.fromisoformat(data["started"]) if data.get("started") else dt.datetime.now()
    opts = data.get("options", {})
    return {"du_an": settings.get("project_name") or data.get("project") or "DuAn",
            "module": mods[0] if len(mods) == 1 else ("NhieuModule" if mods else "TatCa"),
            "ngay": started.strftime("%Y-%m-%d"), "gio": started.strftime("%H%M"),
            "moi_truong": data.get("env") or "", "lan_chay": str(run_no),
            "nguoi_test": opts.get("tester") or settings.get("tester") or "", "trinh_duyet": opts.get("browser", "")}


def build_name(pattern, ctx):
    pattern = (pattern or "").strip() or "KetQua_{du_an}_{ngay}_{gio}"

    def rep(m):
        k = m.group(1).strip()
        return str(ctx.get(k, m.group(0)))
    name = re.sub(r"\{([a-z_]+)\}", rep, pattern)
    if name.lower().endswith(".xlsx"):
        name = name[:-5]
    name = re.sub(BAD_CHARS, "_", name)
    name = re.sub(r"\s+", " ", name).strip(" .")
    name = re.sub(r"_+", "_", name)
    if not name:
        name = "KetQua"
    if name.upper() in ("CON", "PRN", "AUX", "NUL") or re.match(r"^(COM|LPT)\d$", name.upper()):
        name = "_" + name
    return name[:150] + ".xlsx"


def unique_path(path):
    if not os.path.exists(path):
        return path
    stem, ext = os.path.splitext(path)
    i = 2
    while os.path.exists("%s_(%d)%s" % (stem, i, ext)):
        i += 1
    return "%s_(%d)%s" % (stem, i, ext)


def _fill(cell, status):
    if status in STATUS_FILL:
        cell.fill = PatternFill("solid", fgColor=STATUS_FILL[status])
        cell.font = Font(bold=True, color=STATUS_FONT[status])


def scenario_summary(data):
    by_tc = {}
    for r in data["results"]:
        prev = by_tc.get(r["key"])
        by_tc[r["key"]] = r["status"] if not prev else _worse(prev, r["status"])
    out = {}
    for s in data.get("scenarios", []):
        sts = [by_tc[k] for k in s.get("tcs", []) if k in by_tc]
        c = {k: sts.count(k) for k in ("Pass", "Fail", "Blocked", "Chưa test")}
        total = len(sts)
        if not total:
            concl = "Không chạy lần này" if s.get("tcs") else "Chưa có TC"
        elif c["Fail"]:
            concl = "Fail"
        elif c["Blocked"]:
            concl = "Blocked"
        elif c["Chưa test"]:
            concl = "Chưa test xong"
        else:
            concl = "Pass"
        out[s["id"]] = {"tong_tc": total, "so_pass": c["Pass"], "so_fail": c["Fail"], "so_blocked": c["Blocked"],
                        "so_chua_test": c["Chưa test"], "ty_le": round(c["Pass"] * 100.0 / total, 1) if total else "",
                        "ket_luan": concl}
    return out


def _worse(a, b):
    order = ["Fail", "Blocked", "Chưa test", "Pass", "Skip"]
    return a if order.index(a) <= order.index(b) else b


def _tc_value(key, r, data, rnd="1"):
    f = r["fields"]
    if key == "status":
        return r["status"]
    if key in ("kq_lan1", "kq_lan2", "kq_lan3"):
        return r["status"] if key == "kq_lan" + rnd else f.get(key, "")
    if key == "cach_test":
        return r["mode"]
    if key == "thoi_gian":
        return "%gs" % r["duration"]
    if key == "nguoi_test":
        return r.get("tester") or f.get("nguoi_test", "")
    if key == "ngay_test":
        return (r.get("start") or "")[:10]
    if key == "kich_thuoc":
        return r.get("size") or f.get("kich_thuoc", "")
    if key == "trinh_duyet":
        return r.get("browser") or f.get("trinh_duyet", "")
    if key == "moi_truong":
        return r.get("env") or f.get("moi_truong", "")
    if key == "ghi_chu":
        return "; ".join(x for x in [f.get("ghi_chu", ""), r.get("note", "")] if x)
    return f.get(key, "")


def _step_value(key, s):
    if key == "actual":
        return s.get("actual", "")
    if key == "status_buoc":
        return s.get("status", "")
    if key == "step_no":
        return s.get("no", "")
    if key == "buoc" and s.get("no") == "Duyệt":
        return s["fields"].get("buoc", "")
    return s["fields"].get(key, "")


def export_excel(data, profile, out_path, run_dir, embed_images=False, keep_extra=True, opts=None):
    """Ghi file Excel kết quả. Ảnh được chép vào thư mục '<tên file>_anh' cạnh file Excel."""
    folder = os.path.dirname(out_path)
    stem = os.path.splitext(os.path.basename(out_path))[0]
    img_folder_name = stem + "_anh"
    img_folder = os.path.join(folder, img_folder_name)
    copied = {}

    def img(rel):
        if rel in copied:
            return copied[rel]
        src = os.path.join(run_dir, rel)
        if not os.path.isfile(src):
            return None
        os.makedirs(img_folder, exist_ok=True)
        dst = os.path.join(img_folder, os.path.basename(rel))
        shutil.copyfile(src, dst)
        copied[rel] = (img_folder_name + "/" + os.path.basename(rel), dst)
        return copied[rel]

    plan_only = bool((opts or {}).get("plan_only"))   # file test case sinh TRƯỚC khi chạy: chưa có kết quả
    wb = openpyxl.Workbook()
    blank = wb.active
    if not plan_only:
        _summary_sheet(blank, data)

    # ---- Scenario
    ts_cols = X.enabled_columns(profile, "scenario")
    if data.get("scenarios") and ts_cols:
        ws = wb.create_sheet("Scenario")
        X.style_header(ws, ts_cols)
        summ = {} if plan_only else scenario_summary(data)
        started = data.get("started", "")[:10]
        for ri, s in enumerate(data["scenarios"], start=2):
            sm = summ.get(s["id"], {})
            for ci, c in enumerate(ts_cols, start=1):
                k = c["key"]
                if k in sm:
                    v = sm[k]
                elif k == "ngay_chay":
                    v = _vn_date(started)
                elif k == "nguoi_chay":
                    v = data.get("options", {}).get("tester", "")
                else:
                    v = s["fields"].get(k, "")
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.alignment = X.WRAP
                cell.border = X.BORDER
                if k == "ket_luan":
                    _fill(cell, {"Chưa test xong": "Chưa test"}.get(v, v))
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(ts_cols)), max(2, len(data["scenarios"]) + 1))

    # ---- Test case
    tc_cols = X.enabled_columns(profile, "testcase")
    extra_heads = []
    if keep_extra:
        for m in data.get("headers", {}).get("testcase", []):
            if not m["key"] and m["header"] not in extra_heads:
                extra_heads.append(m["header"])
    ws = wb.create_sheet("TestCase")
    all_cols = tc_cols + [{"key": "__x__" + h, "label": h, "group": "Khác", "level": "tc", "result": False}
                          for h in extra_heads]
    X.style_header(ws, all_cols)
    layout = (profile.get("layout") or "rows")
    sections = bool(profile.get("sections"))
    rnd = str((opts or {}).get("round") or "1")

    def put_img(cell, rels, ci, r_i):
        first = img(rels[0]) if rels else None
        if not first:
            return
        for extra in rels[1:]:
            img(extra)
        cell.value = "Xem ảnh" + (" (%d)" % len(rels) if len(rels) > 1 else "")
        cell.hyperlink = first[0]
        cell.font = Font(color="0563C1", underline="single")
        if embed_images:
            try:
                from openpyxl.drawing.image import Image as XLImage
                pic = XLImage(first[1])
                ratio = 170.0 / max(1, pic.width)
                pic.width, pic.height = 170, int(pic.height * ratio)
                ws.add_image(pic, cell.coordinate)
                ws.row_dimensions[r_i].height = max(ws.row_dimensions[r_i].height or 15, pic.height * 0.75 + 4)
                ws.column_dimensions[get_column_letter(ci)].width = max(
                    ws.column_dimensions[get_column_letter(ci)].width or 10, 26)
            except Exception:
                pass

    def section_row(r_i, text):
        for ci in range(1, len(all_cols) + 1):
            c = ws.cell(row=r_i, column=ci, value=text if ci == 1 else None)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
            c.font = Font(bold=True, color="1F4E79")
            c.border = X.BORDER
        return r_i + 1

    r_i = 2
    band = [PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="EEF3FA")]
    last_sec = None
    scn_by_id = {s["id"]: s for s in data.get("scenarios", [])}
    for n, r in enumerate(data["results"]):
        if sections:
            scn = scn_by_id.get(r.get("scenario_id") or "")
            if scn:       # dòng nhóm = TEST SCENARIO mà các test case bên dưới thuộc về
                sec = r["scenario_id"]
                text = "%s — %s%s" % (sec, scn["fields"].get("scenario", ""), (" (%s)" % scn["fields"]["nhom"])
                                       if scn["fields"].get("nhom") else "")
            else:
                sec = (r["fields"].get("loai_test") or "").strip() or "KHÁC"
                text = SECTION_NAME.get(sec.lower(), sec.upper())
            if sec != last_sec:
                r_i = section_row(r_i, text)
                last_sec = sec
        steps = r["steps"] or [{"no": "", "fields": {}, "status": "", "actual": r.get("note", ""), "images": []}]
        if layout == "merged":
            rows = [None]          # 1 dòng / TC, các bước gộp vào 1 ô
        else:
            rows = steps
        for si, s in enumerate(rows):
            for ci, c in enumerate(all_cols, start=1):
                k = c["key"]
                if layout == "merged":
                    v = _merged_value(k, r, data, steps, rnd)
                elif k.startswith("__x__"):
                    h = k[5:]
                    v = (r.get("extra", {}).get(h, "") if si == 0 else "") or s.get("extra", {}).get(h, "")
                elif c["level"] == "step":
                    v = _step_value(k, s)
                elif si == 0:
                    v = _tc_value(k, r, data, rnd)
                    if k == "tc_id" and r.get("label") and r["label"] != r["tc_id"]:
                        v = r["label"]
                else:
                    v = ""
                cell = ws.cell(row=r_i, column=ci, value=v if k != "anh" else None)
                cell.alignment = X.WRAP
                cell.border = X.BORDER
                cell.fill = band[n % 2]
                if k in ("status", "status_buoc", "kq_lan1", "kq_lan2", "kq_lan3") and v:
                    _fill(cell, v)
                if k == "anh":
                    if layout == "merged":
                        bad = [x for x in steps if x.get("status") in ("Fail", "Blocked") and x.get("images")]
                        rels = (bad[0]["images"] if bad else next((x["images"] for x in steps if x.get("images")), []))
                    else:
                        rels = s.get("images") or []
                    put_img(cell, rels, ci, r_i)
            r_i += 1
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(all_cols)), max(2, r_i - 1))

    # ---- Mỗi loại test một bảng riêng (ma trận Validation, bảng quyết định, ma trận phân quyền, chuyển trạng thái…)
    if profile.get("matrices", True) and not plan_only:
        _matrix_sheets(wb, data, img)
    if plan_only:
        wb.remove(blank)
        ws.freeze_panes = "B2"
        wb.move_sheet(ws, offset=-wb.index(ws))     # sheet TestCase đứng đầu
    wb.active = 0
    try:
        wb.save(out_path)
    except PermissionError:
        raise X.ExcelError("Không ghi được file — file đang mở trong Excel? Hãy đóng file hoặc đặt tên khác.\n%s"
                           % out_path)
    return {"path": out_path, "images": len(copied), "image_folder": img_folder if copied else ""}


def _vn_date(iso):
    try:
        return dt.date.fromisoformat(iso[:10]).strftime("%d/%m/%Y")
    except Exception:
        return iso


def counts(data):
    c = {"Pass": 0, "Fail": 0, "Blocked": 0, "Chưa test": 0}
    for r in data["results"]:
        c[r["status"]] = c.get(r["status"], 0) + 1
    return c


def _summary_sheet(ws, data):
    ws.title = "TongQuan"
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 14
    ws["A1"] = "BÁO CÁO KẾT QUẢ TEST"
    ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
    opts = data.get("options", {})
    started = data.get("started") or ""
    finished = data.get("finished") or ""
    dur = ""
    try:
        d = dt.datetime.fromisoformat(finished) - dt.datetime.fromisoformat(started)
        dur = "%d phút %d giây" % (d.seconds // 60, d.seconds % 60)
    except Exception:
        pass
    info = [("Dự án / File test", data.get("project", "")), ("Môi trường", data.get("env", "")),
            ("Trình duyệt", opts.get("browser", "")), ("Người test", opts.get("tester", "")),
            ("Bắt đầu", started.replace("T", " ")), ("Kết thúc", finished.replace("T", " ")), ("Thời gian", dur),
            ("Trạng thái lần chạy", data.get("state", ""))]
    r = 3
    for k, v in info:
        ws.cell(row=r, column=1, value=k).font = Font(bold=True)
        ws.cell(row=r, column=2, value=v)
        r += 1
    r += 1
    c = counts(data)
    total = sum(c.values())
    ws.cell(row=r, column=1, value="Kết quả").font = Font(bold=True, size=12)
    ws.cell(row=r, column=2, value="Số lượt chạy TC").font = Font(bold=True)
    ws.cell(row=r, column=3, value="Tỷ lệ").font = Font(bold=True)
    r0 = r + 1
    for k in ("Pass", "Fail", "Blocked", "Chưa test"):
        r += 1
        a = ws.cell(row=r, column=1, value=k)
        _fill(a, k)
        ws.cell(row=r, column=2, value=c.get(k, 0))
        ws.cell(row=r, column=3, value=("%.1f%%" % (c.get(k, 0) * 100.0 / total)) if total else "")
    r += 1
    ws.cell(row=r, column=1, value="Tổng").font = Font(bold=True)
    ws.cell(row=r, column=2, value=total).font = Font(bold=True)
    if total:
        ch = BarChart()
        ch.type = "bar"
        ch.title = "Kết quả"
        ch.legend = None
        ch.add_data(Reference(ws, min_col=2, min_row=r0, max_row=r0 + 3), titles_from_data=False)
        ch.set_categories(Reference(ws, min_col=1, min_row=r0, max_row=r0 + 3))
        ch.height, ch.width = 6, 12
        ws.add_chart(ch, "E3")
    # theo cách test và loại test
    r += 2
    ws.cell(row=r, column=1, value="Theo cách test").font = Font(bold=True, size=12)
    for mode in ("Tự động", "Bán tự động", "Tay"):
        items = [x for x in data["results"] if x["mode"] == mode]
        if not items:
            continue
        r += 1
        ws.cell(row=r, column=1, value=mode)
        ws.cell(row=r, column=2, value="%d lượt — Pass %d, Fail %d, Blocked %d, Chưa test %d" % (
            len(items), *(sum(1 for x in items if x["status"] == s) for s in ("Pass", "Fail", "Blocked",
                                                                               "Chưa test"))))
    types = {}
    for x in data["results"]:
        types.setdefault(x["fields"].get("loai_test") or "(không ghi)", []).append(x["status"])
    if types:
        r += 2
        ws.cell(row=r, column=1, value="Theo loại test").font = Font(bold=True, size=12)
        for t, sts in sorted(types.items()):
            r += 1
            ws.cell(row=r, column=1, value=t)
            ws.cell(row=r, column=2, value="%d lượt — Pass %d, Fail %d" % (len(sts), sts.count("Pass"),
                                                                          sts.count("Fail")))
    bad = [x for x in data["results"] if x["status"] != "Pass"]
    if bad:
        r += 2
        ws.cell(row=r, column=1, value="Cần chú ý").font = Font(bold=True, size=12, color="9C0006")
        for x in bad:
            r += 1
            ws.cell(row=r, column=1, value=x["label"])
            st = ws.cell(row=r, column=3, value=x["status"])
            _fill(st, x["status"])
            why = x.get("note") or next((s["actual"] for s in x["steps"] if s["status"] in ("Fail", "Blocked")), "")
            if x["status"] == "Chưa test" and not why:
                why = "Case %s chưa được kiểm tra" % x["mode"].lower()
            cell = ws.cell(row=r, column=2, value="%s — %s" % (x["title"], why) if why else x["title"])
            cell.alignment = X.WRAP


# ----------------------------------------------------------------------------- HTML
def export_html(data, out_path, image_folder_name):
    c = counts(data)
    total = sum(c.values()) or 1
    e = html.escape
    rows = []
    for r in data["results"]:
        steps = []
        for s in r["steps"]:
            imgs = "".join('<a href="%s/%s" target="_blank"><img src="%s/%s" loading="lazy"></a>' % (
                e(image_folder_name), e(os.path.basename(i)), e(image_folder_name), e(os.path.basename(i)))
                for i in s.get("images", []))
            steps.append("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td><td><b class='st %s'>%s</b></td>"
                         "<td>%s</td></tr>" % (e(str(s.get("no", ""))), e(s["fields"].get("buoc", "")),
                                               e(s["fields"].get("expected", "")), e(s.get("actual", "")),
                                               _cls(s.get("status")), e(s.get("status", "")), imgs))
        rows.append("""<details class="tc %s" %s><summary><b class="st %s">%s</b> <span class="id">%s</span> %s
            <span class="mode">%s</span><span class="dur">%gs</span></summary>
            %s<table><tr><th>#</th><th>Bước</th><th>Mong đợi</th><th>Thực tế</th><th>KQ</th><th>Ảnh</th></tr>%s
            </table></details>""" % (_cls(r["status"]), "" if r["status"] == "Pass" else "open", _cls(r["status"]),
                                     e(r["status"]), e(r["label"]), e(r["title"]), e(r["mode"]), r["duration"],
                                     "<p class='note'>%s</p>" % e(r["note"]) if r.get("note") else "",
                                     "".join(steps)))
    bars = "".join("<div class='bar %s' style='width:%.2f%%' title='%s: %d'></div>" % (_cls(k), v * 100.0 / total,
                                                                                    k, v) for k, v in c.items() if v)
    page = """<!doctype html><html lang="vi"><head><meta charset="utf-8"><meta name="viewport"
content="width=device-width,initial-scale=1"><title>Báo cáo test %s</title><style>
body{font:14px/1.5 "Segoe UI",Arial,sans-serif;margin:0;background:#f4f6f9;color:#1f2937}
header{background:#1f4e79;color:#fff;padding:16px 24px}h1{margin:0;font-size:20px}
main{padding:16px 24px;max-width:1300px;margin:auto}.sum{display:flex;gap:12px;flex-wrap:wrap;margin:8px 0 16px}
.card{background:#fff;border-radius:8px;padding:10px 16px;box-shadow:0 1px 3px #0002;min-width:110px}
.card b{font-size:22px;display:block}.bars{display:flex;height:14px;border-radius:7px;overflow:hidden;margin:6px 0 18px}
.bar.pass{background:#22c55e}.bar.fail{background:#ef4444}.bar.blocked{background:#eab308}.bar.chua{background:#f97316}
details.tc{background:#fff;border-radius:8px;margin:8px 0;box-shadow:0 1px 3px #0002;border-left:5px solid #ccc}
details.tc.pass{border-color:#22c55e}details.tc.fail{border-color:#ef4444}details.tc.blocked{border-color:#eab308}
details.tc.chua{border-color:#f97316}summary{padding:10px 14px;cursor:pointer}.id{font-weight:600;margin:0 6px}
.mode,.dur{float:right;color:#6b7280;margin-left:12px;font-size:12px}table{border-collapse:collapse;width:100%%}
td,th{border-top:1px solid #e5e7eb;padding:6px 10px;text-align:left;vertical-align:top}th{background:#f9fafb}
.st{padding:1px 8px;border-radius:10px;font-size:12px}.st.pass{background:#dcfce7;color:#166534}
.st.fail{background:#fee2e2;color:#991b1b}.st.blocked{background:#fef9c3;color:#854d0e}
.st.chua{background:#ffedd5;color:#9a3412}.st.skip{background:#e5e7eb;color:#374151}
td img{max-width:180px;max-height:120px;border:1px solid #ddd;margin:2px;border-radius:4px}
.note{margin:0 14px 8px;color:#9a3412}.f button{border:1px solid #cbd5e1;background:#fff;border-radius:6px;
padding:4px 10px;cursor:pointer;margin-right:4px}.f button.on{background:#1f4e79;color:#fff}
</style></head><body><header><h1>Báo cáo kết quả test — %s</h1><div>Môi trường: %s · Trình duyệt: %s ·
Người test: %s · %s</div></header><main><div class="sum">%s</div><div class="bars">%s</div>
<div class="f">Lọc: <button class="on" data-f="">Tất cả</button><button data-f="fail">Fail</button>
<button data-f="blocked">Blocked</button><button data-f="chua">Chưa test</button><button data-f="pass">Pass</button></div>
%s</main><script>document.querySelectorAll('.f button').forEach(b=>b.onclick=()=>{
document.querySelectorAll('.f button').forEach(x=>x.classList.toggle('on',x===b));
document.querySelectorAll('details.tc').forEach(d=>d.style.display=!b.dataset.f||d.classList.contains(b.dataset.f)?'':'none')})
</script></body></html>""" % (e(data.get("project", "")), e(data.get("project", "")), e(data.get("env", "")),
                              e(data.get("options", {}).get("browser", "")), e(data.get("options", {}).get("tester", "")),
                              e((data.get("started") or "").replace("T", " ")),
                              "".join("<div class='card'><b>%d</b>%s</div>" % (v, e(k)) for k, v in
                                      [("Tổng", sum(c.values()))] + list(c.items())),
                              bars, "".join(rows))
    with open(out_path, "w", encoding="utf-8") as f:
        f.write(page)
    return out_path


def _cls(status):
    return {"Pass": "pass", "Fail": "fail", "Blocked": "blocked", "Chưa test": "chua", "Skip": "skip"}.get(status, "")


# ----------------------------------------------------------------------------- ghi ngược vào bản sao file gốc
def write_back(data, src_file, out_path, rnd="1"):
    """Chép file test case gốc rồi điền Status / Actual result / Người test / Ngày test vào đúng dòng.
    File gốc KHÔNG bị sửa."""
    if not src_file or not os.path.isfile(src_file):
        raise X.ExcelError("Không thấy file gốc để ghi kết quả: %s" % src_file)
    shutil.copyfile(src_file, out_path)
    wb = openpyxl.load_workbook(out_path)
    ws = None
    hr = None
    mapping = None
    for w in wb.worksheets:
        k = X.sheet_kind_by_name(w.title) or X.guess_kind(w)
        if k == "testcase":
            hr, mapping = X.find_header(w, "testcase")
            if hr:
                ws = w
                break
    if ws is None:
        raise X.ExcelError("Không tìm thấy sheet Test Case trong file gốc.")
    col = {m["key"]: m["col"] + 1 for m in mapping if m["key"]}
    last = max([m["col"] + 1 for m in mapping] + [ws.max_column])
    for key in ("status", "actual", "nguoi_test", "ngay_test"):
        if key not in col:
            last += 1
            col[key] = last
            cell = ws.cell(row=hr, column=last, value=C.BY_KEY["testcase"][key]["label"])
            cell.font = Font(bold=True)
    merged = {}
    by_key = {}
    for r in data["results"]:
        by_key.setdefault(r["key"], []).append(r)
    for key, items in by_key.items():
        status = items[0]["status"]
        for it in items[1:]:
            status = _worse(status, it["status"])
        row = items[0].get("row")
        if not row:
            continue
        c = ws.cell(row=row, column=col["status"], value=status)
        _fill(c, status)
        if "kq_lan" + rnd in col:          # file có cột "Kết quả test (Lần N)" → điền đúng lần đang test
            _fill(ws.cell(row=row, column=col["kq_lan" + rnd], value=status), status)
        ws.cell(row=row, column=col["nguoi_test"], value=items[0].get("tester", ""))
        ws.cell(row=row, column=col["ngay_test"], value=(items[0].get("start") or "")[:10])
        for si, s in enumerate(items[0]["steps"]):
            srow = s.get("row")
            if not srow:
                continue
            for it in items:
                if si < len(it["steps"]):
                    st = it["steps"][si]
                    pre = (it["size"] + ": ") if len(items) > 1 and it.get("size") else ""
                    no = ("%s. " % st.get("no")) if items[0].get("steps_merged") else ""
                    merged.setdefault((srow, col["actual"]), []).append(
                        "%s%s[%s] %s" % (pre, no, st.get("status", ""), st.get("actual", "")))
    for (r, cidx), v in merged.items():
        cell = ws.cell(row=r, column=cidx, value="\n".join(v))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    try:
        wb.save(out_path)
    except PermissionError:
        raise X.ExcelError("Không ghi được file (đang mở?): %s" % out_path)
    return out_path


SECTION_NAME = {"validation": "VALIDATION — KIỂM TRA DỮ LIỆU NHẬP", "functional": "FUNCTIONAL — CHỨC NĂNG",
                "integration": "INTEGRATION — TÍCH HỢP FRONTEND / BACKEND",
                "business rule": "BUSINESS RULE — LUẬT NGHIỆP VỤ", "permission": "PERMISSION — PHÂN QUYỀN",
                "state": "STATE — TRẠNG THÁI", "usability": "USABILITY — DỄ SỬ DỤNG", "ui": "UI — GIAO DIỆN",
                "responsive": "RESPONSIVE — NHIỀU KÍCH THƯỚC MÀN HÌNH", "performance": "PERFORMANCE — HIỆU NĂNG",
                "security": "SECURITY — BẢO MẬT", "accessibility": "ACCESSIBILITY — KHẢ NĂNG TRUY CẬP",
                "compatibility": "COMPATIBILITY — TƯƠNG THÍCH TRÌNH DUYỆT"}


def _numbered(items):
    items = [x for x in items if x[1]]
    if not items:
        return ""
    if len(items) == 1 and len(items) == 1 and not items[0][0]:
        return items[0][1]
    return "\n".join("%s. %s" % (no, t) for no, t in items)


def _merged_value(key, r, data, steps, rnd):
    """Kiểu file của tester: 1 dòng / test case, các bước gộp vào 1 ô có đánh số."""
    if key == "buoc":
        return _numbered([(s.get("no"), s["fields"].get("buoc", "")) for s in steps])
    if key == "expected":
        return _numbered([(s.get("no"), s["fields"].get("expected", "")) for s in steps])
    if key == "actual":
        bad = [s for s in steps if s.get("status") not in ("Pass", "Skip", "")]
        pick = bad or steps
        return _numbered([(s.get("no"), (("[%s] " % s["status"]) if s.get("status") not in ("Pass", "") else "")
                           + (s.get("actual") or "")) for s in pick])
    if key == "du_lieu_buoc":
        return "\n".join(x for x in (s["fields"].get("du_lieu_buoc", "") for s in steps) if x)
    if key == "status_buoc":
        return r["status"]
    if key == "step_no":
        return len(steps)
    if key in ("hanh_dong", "doi_tuong", "gia_tri"):
        return "\n".join(x for x in (s["fields"].get(key, "") for s in steps) if x)
    if key.startswith("__x__"):
        return r.get("extra", {}).get(key[5:], "")
    v = _tc_value(key, r, data, rnd)
    if key == "tc_id" and r.get("label") and r["label"] != r["tc_id"]:
        v = r["label"]
    return v


MATRIX = {
    "validation": ("Ma trận Validation", [("screen", "Màn hình"), ("form", "Form"), ("field", "Trường"),
                                          ("ftype", "Kiểu dữ liệu"), ("required", "Bắt buộc"),
                                          ("rule", "Luật kiểm tra"), ("input", "Giá trị nhập"),
                                          ("expect", "Kỳ vọng"), ("exp_msg", "Thông báo mong đợi"),
                                          ("act_msg", "Kết quả thực tế")]),
    "business": ("Bảng quyết định (Business rule)", [("ac", "AC liên quan"), ("field", "Trường"),
                                                    ("rule", "Luật nghiệp vụ"), ("input", "Điều kiện / Dữ liệu nhập"),
                                                    ("expect", "Kỳ vọng"),
                                                    ("exp_msg", "Thông báo / Giá trị mong đợi"),
                                                    ("act_msg", "Kết quả thực tế")]),
    "permission": ("Ma trận phân quyền", [("screen", "Màn hình"), ("where", "Vị trí"), ("action", "Chức năng / Nút"),
                                          ("role", "Vai trò"), ("expect", "Kỳ vọng"), ("actual", "Thực tế")]),
    "state": ("Bảng chuyển trạng thái", [("screen", "Màn hình"), ("object", "Đối tượng"),
                                         ("from", "Trạng thái ban đầu"), ("action", "Hành động / Sự kiện"),
                                         ("expect", "Trạng thái mong đợi"), ("actual", "Trạng thái thực tế")]),
    "integration": ("Kiểm tra tích hợp", [("screen", "Màn hình"), ("component", "Thành phần"), ("api", "API"),
                                          ("http", "HTTP"), ("ui_data", "Dữ liệu hiển thị")]),
    "usability": ("Kiểm tra Usability", [("screen", "Màn hình"), ("size", "Thiết bị"), ("criterion", "Tiêu chí"),
                                         ("element", "Phần tử"), ("issue", "Vấn đề")]),
}
GENERIC = {"functional": "Checklist Functional", "ui": "Checklist UI", "responsive": "Checklist Responsive",
           "performance": "Checklist Performance", "security": "Checklist Security",
           "accessibility": "Checklist Accessibility", "compat": "Checklist Compatibility"}


def _matrix_sheets(wb, data, img):
    by_type = {}
    for r in data["results"]:
        if r.get("type"):
            by_type.setdefault(r["type"], []).append(r)
    order = ["validation", "functional", "integration", "business", "permission", "state", "usability", "ui",
             "responsive", "performance", "security", "accessibility", "compat"]
    for t in order:
        res = by_type.get(t)
        if not res:
            continue
        if t in MATRIX:
            title, cols = MATRIX[t]
            rows = [(r, s) for r in res for s in r["steps"] if s.get("m")]
            if not rows:
                continue
            cols = cols + [("_actual", "Ghi chú thực tế"), ("_status", "Kết quả"), ("_img", "Ảnh"), ("_tc", "Mã TC")]
        else:
            title = GENERIC[t]
            rows = [(r, s) for r in res for s in r["steps"]]
            cols = [("_screen", "Màn hình"), ("_title", "Test case"), ("_buoc", "Hạng mục kiểm tra"),
                    ("_exp", "Kỳ vọng"), ("_actual", "Thực tế"), ("_status", "Kết quả"), ("_img", "Ảnh"),
                    ("_tc", "Mã TC")]
        ws = wb.create_sheet(title[:31])
        head = [{"key": k, "label": lab, "result": k in ("_status",), "group": ""} for k, lab in cols]
        X.style_header(ws, head)
        for ci, (k, _) in enumerate(cols, start=1):
            ws.column_dimensions[get_column_letter(ci)].width = {"_tc": 11, "_status": 11, "_img": 11, "required": 9,
                                                                 "http": 8, "ftype": 10, "role": 12}.get(k, 24)
        for ri, (r, s) in enumerate(rows, start=2):
            m = s.get("m") or {}
            for ci, (k, _) in enumerate(cols, start=1):
                if k == "_status":
                    v = s.get("status", "")
                elif k == "_tc":
                    v = r["tc_id"]
                elif k == "_actual":
                    v = s.get("actual", "")
                elif k == "_screen":
                    v = r["fields"].get("module", "")
                elif k == "_title":
                    v = r["title"]
                elif k == "_buoc":
                    v = s["fields"].get("buoc", "")
                elif k == "_exp":
                    v = s["fields"].get("expected", "")
                elif k == "_img":
                    v = None
                else:
                    v = m.get(k, "")
                cell = ws.cell(row=ri, column=ci, value=v)
                cell.alignment = X.WRAP
                cell.border = X.BORDER
                if k == "_status" and v:
                    _fill(cell, v)
                if k in ("expect", "actual") and t == "permission" and v:
                    cell.font = Font(bold=True, color="006100" if v in ("Được phép", "Thấy", "Vào được") else "9C0006")
                if k == "_img" and s.get("images"):
                    got = img(s["images"][0])
                    if got:
                        cell.value = "Xem ảnh"
                        cell.hyperlink = got[0]
                        cell.font = Font(color="0563C1", underline="single")
        ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(cols)), max(2, len(rows) + 1))


# ----------------------------------------------------------------------------- file TEST SCENARIO theo file mẫu tester
TS_BLUE = "1F4E78"
TS_GROUP_FILL = "DDEBF7"
PRIO_FILL = {"high": "F8CBAD", "critical": "F8CBAD", "medium": "FFE699", "low": "E2EFDA"}
_S9 = Side(style="thin", color="999999")
B9 = Border(left=_S9, right=_S9, top=_S9, bottom=_S9)
TS_WIDTH = {"ts_id": 14, "nhom": 13, "scenario": 95, "ac": 20, "priority": 10, "module": 24, "requirement": 22,
            "luong": 40, "cach_test": 13, "ghi_chu": 30}
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)


def luong_title(no, name, mhs):
    return "LUỒNG %d – %s%s" % (no, name, (" (%s)" % ", ".join(mhs)) if mhs else "")


def _head(ws, row, labels):
    for ci, lab in enumerate(labels, start=1):
        c = ws.cell(row=row, column=ci, value=lab)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TS_BLUE)
        c.border = B9
        c.alignment = CENTER


def ts_value(key, r):
    if key == "luong":
        return r.get("group", "")
    if key == "cach_test":
        return r.get("mode", "")
    if key == "module" and r.get("mh"):
        return "%s – %s" % (r["mh"], r.get("module", ""))
    return r.get("fields", {}).get(key, "")


def scenario_groups(rows):
    """[(tiêu đề LUỒNG, [dòng…])] — luồng đánh số lại liên tục theo những luồng có scenario."""
    out = []
    for r in rows:
        if not out or out[-1][0] != r.get("luong"):
            out.append([r.get("luong"), []])
        out[-1][1].append(r)
    res = []
    for i, (lu, items) in enumerate(out, start=1):
        mhs = list(dict.fromkeys(x["mh"] for x in items if x.get("mh")))
        title = luong_title(i, items[0].get("luong_name") or "KHÁC", mhs)
        for x in items:
            x["group"] = title
        res.append((title, items))
    return res


def export_scenarios(path, rows, cols, meta):
    """Sheet 'Test Scenario' đúng bố cục file mẫu: dòng 1 tiêu đề, dòng 2 ghi chú, dòng 4 header, dòng LUỒNG gộp ô,
    freeze A5; kèm sheet 'Ma tran truy vet' và 'Tong hop' như file mẫu."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Scenario"
    n = len(cols)
    ws["A1"] = "TEST SCENARIO – MÀN HÌNH %s" % (meta.get("title") or "").upper()
    ws["A1"].font = Font(bold=True, size=14, color=TS_BLUE)
    ws["A2"] = meta.get("note", "")
    ws["A2"].font = Font(italic=True, color="555555")
    _head(ws, 4, [c["label"] for c in cols])
    for ci, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(ci)].width = TS_WIDTH.get(c["key"], 16)
    groups = scenario_groups(rows)
    r = 5
    for title, items in groups:
        for ci in range(1, n + 1):
            c = ws.cell(row=r, column=ci, value=title if ci == 1 else None)
            c.fill = PatternFill("solid", fgColor=TS_GROUP_FILL)
            c.font = Font(bold=True, color=TS_BLUE)
            c.border = B9
        if n > 1:
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        r += 1
        for x in items:
            for ci, c in enumerate(cols, start=1):
                v = ts_value(c["key"], x)
                cell = ws.cell(row=r, column=ci, value=v if v != "" else None)
                cell.border = B9
                cell.font = Font(name="Calibri", size=11)
                cell.alignment = LEFT if c["key"] in ("scenario", "ghi_chu", "luong") else CENTER
                if c["key"] == "priority" and PRIO_FILL.get(str(v).lower()):
                    cell.fill = PatternFill("solid", fgColor=PRIO_FILL[str(v).lower()])
            r += 1
    ws.freeze_panes = "A5"
    _trace_sheet(wb, rows, meta)
    _ts_summary(wb, groups)
    try:
        wb.save(path)
    except PermissionError:
        raise X.ExcelError("Không ghi được file — file đang mở trong Excel? Hãy đóng file hoặc đặt tên khác.\n%s" % path)
    return {"path": path, "rows": len(rows), "groups": len(groups)}


def _trace_sheet(wb, rows, meta):
    ws = wb.create_sheet("Ma tran truy vet")
    for col, w in zip("ABCD", (10, 60, 12, 70)):
        ws.column_dimensions[col].width = w
    ws["A1"] = "MA TRẬN TRUY VẾT – mỗi màn hình / luồng / nhóm test đều có Test Scenario"
    ws["A1"].font = Font(bold=True, size=14, color=TS_BLUE)
    top = Alignment(vertical="top", wrap_text=True)
    sections = []
    scr = [(s["mh"], s["name"]) for s in meta.get("screens", [])]
    sections.append(("Màn hình (MH)", [(mh, nm, [x for x in rows if x.get("mh") == mh]) for mh, nm in scr]))
    lu = list(dict.fromkeys(x["group"] for x in rows))
    sections.append(("Luồng", [("LU_%02d" % (i + 1), g.split("–", 1)[-1].strip(), [x for x in rows if x["group"] == g])
                               for i, g in enumerate(lu)]))
    nh = list(dict.fromkeys(x["fields"].get("nhom", "") for x in rows))
    sections.append(("Nhóm test", [(g, "Nhóm %s" % g, [x for x in rows if x["fields"].get("nhom", "") == g]) for g in nh]))
    r = 3
    for title, items in sections:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, color=TS_BLUE)
        r += 1
        for ci, lab in enumerate(("Mã", "Nội dung", "Số scenario", "Test Scenario ID"), start=1):
            c = ws.cell(row=r, column=ci, value=lab)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor=TS_BLUE)
            c.border = B9
        r += 1
        for code, text, xs in items:
            for ci, v in enumerate((code, text, len(xs), ", ".join(x["fields"]["ts_id"] for x in xs)), start=1):
                c = ws.cell(row=r, column=ci, value=v)
                c.border = B9
                c.alignment = top
            r += 1
        r += 1


def _ts_summary(wb, groups):
    ws = wb.create_sheet("Tong hop")
    for col, w in zip("ABCDEF", (75, 12, 8, 10, 8, 26)):
        ws.column_dimensions[col].width = w
    heads = ("Luồng", "Số scenario", "High", "Medium", "Low", "ID")
    for ci, lab in enumerate(heads, start=1):
        c = ws.cell(row=1, column=ci, value=lab)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TS_BLUE)
        c.border = B9
    mid = Alignment(vertical="center", wrap_text=True)
    tot = [0, 0, 0, 0]
    r = 2
    for title, items in groups:
        pr = [str(x["fields"].get("priority", "")).lower() for x in items]
        vals = [len(items), sum(p in ("high", "critical") for p in pr), pr.count("medium"), pr.count("low")]
        tot = [a + b for a, b in zip(tot, vals)]
        ids = [x["fields"]["ts_id"] for x in items]
        for ci, v in enumerate([title] + vals + [ids[0] if len(ids) == 1 else "%s → %s" % (ids[0], ids[-1])], start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.border = B9
            c.alignment = mid
        r += 1
    for ci, v in enumerate(["TỔNG"] + tot, start=1):
        c = ws.cell(row=r, column=ci, value=v)
        c.font = Font(bold=True)
        c.border = B9
        c.alignment = mid


# ----------------------------------------------------------------------------- ĐIỀN vào file Excel sẵn có
def _cell_txt(v):
    return X.cell_text(v)


def sheet_info(path, profile=None):
    """Các sheet của file + dòng tiêu đề / cột tool nhận ra (cho người dùng chọn sheet để điền)."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except PermissionError:
        raise X.ExcelError("File đang bị khóa (đang mở bằng chương trình khác?). Hãy đóng file rồi thử lại.")
    except Exception as e:
        raise X.ExcelError("Không đọc được file Excel (.xlsx/.xlsm): %s" % e)
    out = []
    for ws in wb.worksheets:
        best = None
        for kind in ("scenario", "testcase"):
            hr, mp = X.find_header(ws, kind, profile)
            if not hr:
                continue
            keys = {m["key"] for m in mp if m["key"]}
            score = len(keys) + (3 if kind == "testcase" and "tc_id" in keys and keys & {"buoc", "expected", "tieu_de"}
                                 else 0) + (3 if kind == "scenario" and {"ts_id", "scenario"} <= keys else 0)
            if not best or score > best[0]:
                best = (score, kind, hr, mp)
        info = {"name": ws.title, "kind": None, "header_row": None, "cols": [], "rows": 0, "groups": 0}
        if best:
            _, kind, hr, mp = best
            data, groups = _data_rows(ws, hr, mp)
            info.update({"kind": kind, "header_row": hr, "rows": len(data), "groups": len(groups),
                         "cols": [{"header": m["header"], "key": m["key"]} for m in mp]})
        out.append(info)
    return out


def _data_rows(ws, hr, mapping):
    """Dòng dữ liệu và dòng nhóm (VD 'LUỒNG 1 – …' gộp ô) bên dưới dòng tiêu đề."""
    last = max([m["col"] + 1 for m in mapping] or [1])
    merged = {(m.min_row, m.min_col) for m in ws.merged_cells.ranges if m.max_col > m.min_col}
    data, groups = [], []
    for r in range(hr + 1, (ws.max_row or hr) + 1):
        vals = [_cell_txt(ws.cell(row=r, column=c).value) for c in range(1, last + 1)]
        filled = [v for v in vals if v]
        if not filled:
            continue
        if (len(filled) == 1 and vals[0] and (last >= 3 or (r, 1) in merged or X.is_section(vals[0]) or
                                              re.match(r"^(LUỒNG|LUONG|NHÓM|PHẦN)\b", vals[0], re.I))):
            groups.append(r)
        else:
            data.append(r)
    return data, groups


def _copy_style(src, dst, fill=True):
    if src is None or not src.has_style:
        return
    dst.font = copy.copy(src.font)
    dst.border = copy.copy(src.border)
    dst.alignment = copy.copy(src.alignment)
    dst.number_format = src.number_format
    dst.protection = copy.copy(src.protection)
    if fill:
        dst.fill = copy.copy(src.fill)


def fill_existing(src, out, sheet, kind, profile, items, mode="append", id_key=None, update_keys=(), rnd="1",
                  renumber="always"):
    """Ghi vào file Excel CÓ SẴN (đã chép sang `out`, trừ khi out == src và người dùng đồng ý ghi đè):
    nhận cột theo TÊN tiêu đề, giữ định dạng (chép kiểu chữ / viền / màu từ dòng có sẵn).
      items : [("group", 'LUỒNG 1 – …') | ("row", {key: giá trị})]
      mode  : 'append' = ghi tiếp dưới dòng cuối (ID nối tiếp file); 'replace' = xóa dữ liệu cũ rồi ghi;
              'update' = tìm dòng có cùng mã (id_key) VÀ cùng tên → chỉ ghi các cột update_keys, không thấy thì thêm dòng
      renumber: 'always' = ID ghi thêm nối tiếp file; 'collision' = chỉ đánh lại khi trùng mã đã có trong file."""
    if not os.path.isfile(src):
        raise X.ExcelError("Không thấy file: %s" % src)
    if os.path.abspath(out) != os.path.abspath(src):
        shutil.copyfile(src, out)
    keep_vba = out.lower().endswith(".xlsm")
    try:
        wb = openpyxl.load_workbook(out, keep_vba=keep_vba)
    except PermissionError:
        raise X.ExcelError("File đang bị khóa (đang mở trong Excel?). Hãy đóng file rồi thử lại.")
    if sheet not in wb.sheetnames:
        raise X.ExcelError("File không có sheet '%s'." % sheet)
    ws = wb[sheet]
    hr, mapping = X.find_header(ws, kind, profile)
    if not hr:
        raise X.ExcelError("Sheet '%s' không có dòng tiêu đề %s mà tool nhận ra (VD: %s)." % (
            sheet, "Test Scenario" if kind == "scenario" else "Test Case",
            "ID, Nhóm, Test Scenario" if kind == "scenario" else "Mã TC, Tên TC, Steps"))
    col = {m["key"]: m["col"] + 1 for m in mapping if m["key"]}
    last_col = max([m["col"] + 1 for m in mapping])
    data, groups = _data_rows(ws, hr, mapping)
    st_row = data[0] if data else None
    st_cells = {c: copy.copy(ws.cell(row=st_row, column=c)._style) if st_row else None for c in range(1, last_col + 1)}
    g_style = copy.copy(ws.cell(row=groups[0], column=1)._style) if groups else None
    prio_col = col.get("priority")
    prio_colored = bool(prio_col and (not st_row or ws.cell(row=st_row, column=prio_col).fill.fill_type))
    id_col = col.get(id_key) if id_key else None
    stats = {"appended": 0, "updated": 0, "groups": 0, "sheet": sheet, "header_row": hr, "path": out,
             "matched": [m["header"] for m in mapping if m["key"]]}

    if data and not groups:      # file không dùng dòng nhóm → không chèn dòng nhóm
        items = [it for it in items if it[0] != "group"]
    title_key = "scenario" if kind == "scenario" else "tieu_de"
    title_col = col.get(title_key)
    have_ids = {_cell_txt(ws.cell(row=r, column=id_col).value) for r in data} if id_col else set()
    clash = any(k == "row" and str(v.get(id_key, "")) in have_ids for k, v in items)
    # ID nối tiếp file (VD file đã có TS_DHTD_043 → dòng mới bắt đầu TS_DHTD_044)
    if mode == "append" and id_col and data and (renumber == "always" or clash):
        nums, prefix, width = [], None, 3
        for r in data:
            m = re.match(r"^(.*?)(\d+)$", _cell_txt(ws.cell(row=r, column=id_col).value))
            if m:
                prefix, width = m.group(1), len(m.group(2))
                nums.append(int(m.group(2)))
        if prefix is not None:
            n = max(nums)
            for kind_, vals in items:
                if kind_ == "row" and vals.get(id_key):
                    n += 1
                    vals[id_key] = "%s%0*d" % (prefix, width, n)
            stats["id_from"] = "%s%0*d" % (prefix, width, max(nums) + 1)
    if mode == "append" and groups:        # file đã có LUỒNG 1…9 → luồng ghi thêm đánh số tiếp (LUỒNG 10…)
        have = [int(m.group(1)) for r in groups for m in [re.match(r"^\s*LU[ỒO]NG\s+(\d+)",
                                                                   _cell_txt(ws.cell(row=r, column=1).value), re.I)] if m]
        if have:
            k = max(have)
            for i, (kind_, vals) in enumerate(items):
                if kind_ == "group":
                    k += 1
                    items[i] = ("group", re.sub(r"^\s*LUỒNG\s+\d+", "LUỒNG %d" % k, vals))

    if mode == "replace":
        for m in list(ws.merged_cells.ranges):
            if m.min_row > hr:
                ws.unmerge_cells(str(m))
        if ws.max_row > hr:
            ws.delete_rows(hr + 1, ws.max_row - hr)
        data, groups = [], []
    index, by_title = {}, {}
    if mode == "update":
        for r in data:
            if id_col:
                index.setdefault(_cell_txt(ws.cell(row=r, column=id_col).value), r)
            if title_col:
                by_title.setdefault(_cell_txt(ws.cell(row=r, column=title_col).value), r)

    def find_row(vals):
        """Cùng mã và cùng tên (nếu file có cột tên) → đúng dòng; không thì tìm theo tên."""
        r = index.get(str(vals.get(id_key, "")))
        t = str(vals.get(title_key, "") or "")
        if r and title_col and t and _cell_txt(ws.cell(row=r, column=title_col).value) != t:
            r = None
        return r or (by_title.get(t) if t else None)
    nxt = max([hr] + data + groups) + 1

    def put(r, key, v, full):
        c = col.get(key)
        if not c:
            return
        cell = ws.cell(row=r, column=c)
        if full and st_cells.get(c) is not None:
            cell._style = copy.copy(st_cells[c])
        cell.value = v if v not in ("",) else None
        if key == "priority" and prio_colored and PRIO_FILL.get(str(v).lower()):
            cell.fill = PatternFill("solid", fgColor=PRIO_FILL[str(v).lower()])
        if key in ("status", "status_buoc", "kq_lan1", "kq_lan2", "kq_lan3") and v in STATUS_FILL:
            _fill(cell, v)
        if full and st_cells.get(c) is None:         # sheet chỉ có tiêu đề → dùng kiểu của file mẫu
            cell.border = B9
            cell.alignment = LEFT if key in ("scenario", "tieu_de", "buoc", "expected", "actual", "precondition") \
                else CENTER

    for kind_, vals in items:
        if kind_ == "group":
            if mode == "update":
                continue
            for c in range(1, last_col + 1):
                cell = ws.cell(row=nxt, column=c, value=vals if c == 1 else None)
                if g_style is not None:
                    cell._style = copy.copy(g_style)
                else:
                    cell.fill = PatternFill("solid", fgColor=TS_GROUP_FILL)
                    cell.font = Font(bold=True, color=TS_BLUE)
                    cell.border = B9
            if last_col > 1:
                ws.merge_cells(start_row=nxt, start_column=1, end_row=nxt, end_column=last_col)
            stats["groups"] += 1
            nxt += 1
            continue
        r = find_row(vals) if mode == "update" else None
        if r:
            for k in update_keys:
                if k in vals:
                    put(r, k, vals[k], False)
            stats["updated"] += 1
            continue
        if st_row:
            for c in range(1, last_col + 1):
                if c not in col.values():
                    ws.cell(row=nxt, column=c)._style = copy.copy(st_cells[c])
        for k, v in vals.items():
            put(nxt, k, v, True)
        stats["appended"] += 1
        nxt += 1
    try:
        wb.save(out)
    except PermissionError:
        raise X.ExcelError("Không ghi được file (đang mở trong Excel?): %s" % out)
    return stats
